#!/usr/bin/env python3
"""Build the ACKE.in weekly scrum status deck (.pptx) from a status JSON file.

Usage: python build_deck.py --data status.json --out deck.pptx
Run generate_assets.py first if assets are missing (make_report.py does this for you).
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any

import theme
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Inches, Pt
from validate_status import validate_status

ROWS_PER_SLIDE = 8
MAX_CARDS = 3
RAG_LABEL = {"green": "ON TRACK", "amber": "AT RISK", "red": "OFF TRACK"}
RAG_TEXT = {"green": "white", "amber": "dark", "red": "white"}
SEVERITY_COLOR = {"low": "green", "medium": "amber", "high": "red"}


# ------------------------------------------------------------------ helpers --------------------


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(theme.PALETTE[name])


def _style(run: Any, size: float, bold: bool, color: str) -> None:
    run.font.name = theme.FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(slide: Slide, x: float, y: float, w: float, h: float, text: str, *, size: float = 18,
             bold: bool = False, color: str = "dark", align: Any = PP_ALIGN.LEFT,
             anchor: Any = MSO_ANCHOR.TOP, name: str | None = None) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    _style(p.add_run(), size, bold, color)
    p.runs[0].text = text
    return box


def add_bullets(slide: Slide, x: float, y: float, w: float, h: float, items: list[str], *,
                size: float = 18, name: str = "Bullets") -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        _style(p.add_run(), size, False, "dark")
        p.runs[0].text = item
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "285750")
        ppr.set("indent", "-285750")
        clr = ppr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": theme.PALETTE["red"]}))
        ppr.append(clr)
        ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "\u2022"}))
    return box


def add_rect(slide: Slide, x: float, y: float, w: float, h: float, fill: str, *,
             shape: Any = MSO_SHAPE.RECTANGLE, name: str = "Shape") -> Any:
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_picture(slide: Slide, path: Path, x: float, y: float, *, height: float | None = None,
                width: float | None = None, name: str) -> Any:
    kwargs: dict[str, Any] = {}
    if height is not None:
        kwargs["height"] = Inches(height)
    if width is not None:
        kwargs["width"] = Inches(width)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)
    pic.name = name
    return pic


def set_notes(slide: Slide, lines: list[str]) -> None:
    if lines:
        slide.notes_slide.notes_text_frame.text = "\n".join(lines)


# ------------------------------------------------------------------ data helpers --------------


def _fmt_date(d: date) -> str:
    return f"{d.day} {d.strftime('%b %Y')}"


def week_range(data: dict[str, Any]) -> str:
    start, end = date.fromisoformat(data["week_start"]), date.fromisoformat(data["week_end"])
    if start.year == end.year and start.month == end.month:
        return f"{start.day} \u2013 {_fmt_date(end)}"
    return f"{_fmt_date(start)} \u2013 {_fmt_date(end)}"


def week_label(data: dict[str, Any]) -> str:
    return f"W{date.fromisoformat(data['week_start']).isocalendar().week:02d}"


def metrics(data: dict[str, Any]) -> dict[str, Any]:
    done = sum(i["points"] for i in data["completed"])
    wip = sum(i["points"] for i in data["in_progress"])
    committed = data.get("committed_points") or (done + wip)
    pct = round(100 * done / committed) if committed else 0
    return {"done": done, "committed": committed, "pct": pct,
            "items_done": len(data["completed"]), "blocked": len(data["blocked"])}


def fmt_num(v: float) -> str:
    return f"{v:g}"


# ------------------------------------------------------------------ slide chrome --------------


class Deck:
    def __init__(self, data: dict[str, Any]) -> None:
        self.data = data
        self.prs: PresentationType = Presentation()
        self.prs.slide_width = Inches(theme.SLIDE_W_IN)
        self.prs.slide_height = Inches(theme.SLIDE_H_IN)
        self.blank = self.prs.slide_layouts[6]
        self.footer = (f"ACKE.in  |  Weekly Scrum Status  |  {data['team']}  |  "
                       f"Sprint {data['sprint']}  |  {week_label(data)}")

    def _slide(self, background: str) -> Slide:
        slide = self.prs.slides.add_slide(self.blank)
        pic = slide.shapes.add_picture(str(theme.ASSETS_DIR / background), 0, 0,
                                       self.prs.slide_width, self.prs.slide_height)
        pic.name = "Background"
        return slide

    def content(self, title: str, icon: str) -> Slide:
        slide = self._slide("bg-content.png")
        add_text(slide, 0.6, 0.4, 11.4, 0.8, title, size=32, bold=True, color="blue",
                 anchor=MSO_ANCHOR.MIDDLE, name="Title")
        add_rect(slide, 0.6, 1.22, 0.9, 0.06, "red", name="Accent")
        add_picture(slide, theme.ICONS_DIR / f"{icon}.png", 12.1, 0.45, height=0.7, name="Icon")
        add_picture(slide, theme.ASSETS_DIR / "acke-logo-mark.png", 0.6, 7.08, height=0.32,
                    name="ACKE Mark")
        add_text(slide, 1.05, 7.08, 9.5, 0.32, self.footer, size=11, anchor=MSO_ANCHOR.MIDDLE,
                 name="Footer")
        add_text(slide, 11.7, 7.08, 1.0, 0.32, str(len(self.prs.slides)), size=11,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name="Page")
        return slide

    def hero(self) -> Slide:
        return self._slide("bg-title.png")

    # ---------------------------------------------------------------- slides ----------------

    def title_slide(self) -> None:
        d = self.data
        slide = self.hero()
        add_picture(slide, theme.ASSETS_DIR / "acke-logo-white.png", 0.8, 0.8, height=0.9,
                    name="ACKE Logo")
        add_text(slide, 0.8, 2.7, 9.5, 1.6, "Weekly Scrum\nStatus Report", size=48, bold=True,
                 color="white", name="Title")
        add_text(slide, 0.8, 4.55, 9.5, 0.5, f"{d['team']}  \u00b7  Sprint {d['sprint']}", size=24,
                 color="light_blue", name="Subtitle")
        add_text(slide, 0.8, 5.15, 9.5, 0.4, f"{week_range(d)}  ({week_label(d)})", size=18,
                 color="white", name="Dates")
        if d.get("prepared_by"):
            add_text(slide, 0.8, 6.5, 9.5, 0.3, f"Prepared by {d['prepared_by']}", size=14,
                     color="white", name="PreparedBy")

    def summary_slide(self) -> None:
        d, m = self.data, metrics(self.data)
        slide = self.content("Executive Summary", "goal")
        rag = d["rag"]
        chip = add_rect(slide, 0.6, 1.55, 2.0, 0.5, rag, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                        name="RAG")
        tf = chip.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        _style(tf.paragraphs[0].add_run(), 14, True, RAG_TEXT[rag])
        tf.paragraphs[0].runs[0].text = RAG_LABEL[rag]
        add_text(slide, 2.85, 1.55, 5.2, 0.5, d["rag_reason"], size=16, color="dark",
                 anchor=MSO_ANCHOR.MIDDLE, name="RAGReason")
        add_bullets(slide, 0.6, 2.45, 7.5, 4.2, d["summary"])
        cards = [("Sprint progress", f"{m['pct']}%"),
                 ("Points done", f"{fmt_num(m['done'])}/{fmt_num(m['committed'])}"),
                 ("Items done", str(m["items_done"])),
                 ("Items blocked", str(m["blocked"]))]
        for i, (label, value) in enumerate(cards):
            x, y = 8.7 + (i % 2) * 2.2, 1.55 + (i // 2) * 2.4
            add_rect(slide, x, y, 2.0, 2.2, "gray", name="Card")
            add_rect(slide, x, y, 2.0, 0.07, "blue", name="CardAccent")
            add_text(slide, x, y + 0.5, 2.0, 0.8, value, size=34, bold=True, color="blue",
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="KPI")
            add_text(slide, x, y + 1.45, 2.0, 0.4, label, size=13, align=PP_ALIGN.CENTER,
                     name="KPILabel")
        set_notes(slide, [d.get("notes", "")] if d.get("notes") else [])

    def goal_slide(self) -> None:
        d, m = self.data, metrics(self.data)
        slide = self.content("Sprint Goal & Burndown", "goal")
        add_text(slide, 0.6, 1.6, 4.4, 0.3, "SPRINT GOAL", size=12, bold=True, color="red")
        add_text(slide, 0.6, 2.0, 4.4, 2.8, d["goal"], size=22, bold=True, color="dark",
                 name="Goal")
        add_text(slide, 0.6, 5.2, 4.4, 0.3, "COMMITTED", size=12, bold=True, color="red")
        add_text(slide, 0.6, 5.55, 4.4, 0.6,
                 f"{fmt_num(m['committed'])} points  \u00b7  {m['pct']}% done", size=20,
                 color="blue", bold=True, name="Committed")
        labels = list(d["burndown"]["labels"])
        remaining = list(d["burndown"]["remaining"])
        n = len(labels)
        ideal = [round(m["committed"] * (1 - i / (n - 1)), 1) for i in range(n)]
        actual = remaining + [None] * (n - len(remaining))
        chart_data = CategoryChartData()
        chart_data.categories = labels
        chart_data.add_series("Ideal", ideal)
        chart_data.add_series("Actual", actual)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(5.4), Inches(1.5),
                                    Inches(7.3), Inches(5.3), chart_data)
        gf.name = "Burndown"
        self._style_chart(gf.chart, [("mid_gray", True), ("red", False)], "Story points remaining")

    def _style_chart(self, chart: Any, series_colors: list[tuple[str, bool]], axis_title: str,
                     *, bars: bool = False) -> None:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.name = theme.FONT
        chart.font.name = theme.FONT
        chart.font.size = Pt(12)
        va = chart.value_axis
        va.has_title = True
        va.axis_title.text_frame.text = axis_title
        run = va.axis_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.name = Pt(12), theme.FONT
        va.major_gridlines.format.line.color.rgb = rgb("gray")
        for series, (color, dashed) in zip(chart.plots[0].series, series_colors, strict=True):
            if bars:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = rgb(color)
            else:
                series.format.line.color.rgb = rgb(color)
                series.format.line.width = Pt(3)
                series.smooth = False
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = rgb(color)
                series.marker.format.line.color.rgb = rgb(color)
                if dashed:
                    series.format.line.dash_style = MSO_LINE.DASH

    def table_slides(self, title: str, icon: str, headers: list[str], widths: list[float],
                     rows: list[tuple[list[str], str | None]], empty_text: str) -> None:
        pages = [rows[i:i + ROWS_PER_SLIDE] for i in range(0, len(rows), ROWS_PER_SLIDE)] or [[]]
        for idx, page in enumerate(pages, start=1):
            suffix = f" ({idx}/{len(pages)})" if len(pages) > 1 else ""
            slide = self.content(f"{title}{suffix}", icon)
            if not page:
                add_text(slide, 0.6, 2.2, 12.0, 0.6, empty_text, size=24, name="Empty")
                continue
            shape = slide.shapes.add_table(len(page) + 1, len(headers), Inches(0.6), Inches(1.6),
                                           Inches(sum(widths)), Inches(0.55 * (len(page) + 1)))
            shape.name = "Table"
            table = shape.table
            table.horz_banding = False
            for col, w in enumerate(widths):
                table.columns[col].width = Inches(w)
            for r in range(len(page) + 1):
                table.rows[r].height = Inches(0.55)
            for c, header in enumerate(headers):
                self._cell(table.cell(0, c), header, "blue", "white", bold=True)
            for r, (cells, _) in enumerate(page, start=1):
                fill = "gray" if r % 2 == 0 else "white"
                for c, text in enumerate(cells):
                    self._cell(table.cell(r, c), text, fill, "dark")
            set_notes(slide, [f"{cells[0]}: {detail}" for cells, detail in page if detail])

    @staticmethod
    def _cell(cell: Any, text: str, fill: str, color: str, *, bold: bool = False) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.12)
        p = cell.text_frame.paragraphs[0]
        _style(p.add_run(), 14, bold, color)
        p.runs[0].text = text

    def completed_slides(self) -> None:
        rows = [([i["id"], i["title"], i["owner"], fmt_num(i["points"])], i.get("detail"))
                for i in self.data["completed"]]
        self.table_slides("Completed", "done", ["ID", "Item", "Owner", "Points"],
                          [1.3, 7.2, 2.4, 1.2], rows, "None this week.")

    def in_progress_slides(self) -> None:
        rows = [([i["id"], i["title"], i["owner"], fmt_num(i["points"]),
                  f"{fmt_num(i['percent'])}%"], i.get("detail"))
                for i in self.data["in_progress"]]
        self.table_slides("In Progress", "progress",
                          ["ID", "Item", "Owner", "Points", "Progress"],
                          [1.3, 6.0, 2.4, 1.2, 1.2], rows, "None this week.")

    def blocked_slide(self) -> None:
        d = self.data
        slide = self.content("Blockers & Risks", "blocked")
        add_text(slide, 0.6, 1.55, 5.9, 0.3, "BLOCKERS", size=12, bold=True, color="red")
        add_text(slide, 6.9, 1.55, 5.9, 0.3, "RISKS", size=12, bold=True, color="red")
        blocked, risks = d["blocked"], d["risks"]
        overflow: list[str] = []
        if not blocked:
            add_text(slide, 0.6, 2.05, 5.9, 0.5, "None this week.", size=20, name="NoBlockers")
        if not risks:
            add_text(slide, 6.9, 2.05, 5.9, 0.5, "None this week.", size=20, name="NoRisks")
        for i, b in enumerate(blocked[:MAX_CARDS]):
            self._card(slide, 0.6, 2.0 + i * 1.6, "red", [
                (f"{b['id']}  \u2013  {b['title']}", 14, True),
                (f"Owner: {b['owner']}", 12, False),
                (f"Ask: {b['ask']}", 12, False)])
        for i, r in enumerate(risks[:MAX_CARDS]):
            owner = f"  \u00b7  Owner: {r['owner']}" if r.get("owner") else ""
            self._card(slide, 6.9, 2.0 + i * 1.6, SEVERITY_COLOR[r["severity"]], [
                (r["risk"], 14, True),
                (f"Severity: {r['severity'].upper()}{owner}", 12, False),
                (f"Mitigation: {r['mitigation']}", 12, False)])
        overflow += [f"Blocker {b['id']}: {b['title']} - ask: {b['ask']}"
                     for b in blocked[MAX_CARDS:]]
        overflow += [f"Risk ({r['severity']}): {r['risk']} - {r['mitigation']}"
                     for r in risks[MAX_CARDS:]]
        notes = overflow + [f"{b['id']}: {b['detail']}" for b in blocked if b.get("detail")]
        set_notes(slide, notes)

    @staticmethod
    def _card(slide: Slide, x: float, y: float, accent: str,
              lines: list[tuple[str, int, bool]]) -> None:
        add_rect(slide, x, y, 5.9, 1.4, "gray", name="Card")
        add_rect(slide, x, y, 0.1, 1.4, accent, name="CardAccent")
        box = slide.shapes.add_textbox(Inches(x + 0.3), Inches(y + 0.12), Inches(5.4), Inches(1.2))
        box.name = "CardText"
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, (text, size, bold) in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(3)
            _style(p.add_run(), size, bold, "dark")
            p.runs[0].text = text

    def velocity_slide(self) -> None:
        vel = self.data["velocity"]
        slide = self.content("Velocity", "velocity")
        chart_data = CategoryChartData()
        chart_data.categories = [v["sprint"] for v in vel]
        chart_data.add_series("Committed", [v["committed"] for v in vel])
        chart_data.add_series("Completed", [v["completed"] for v in vel])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5),
                                    Inches(8.4), Inches(5.3), chart_data)
        gf.name = "VelocityChart"
        self._style_chart(gf.chart, [("light_blue", False), ("blue", False)], "Story points",
                          bars=True)
        total_c, total_d = sum(v["committed"] for v in vel), sum(v["completed"] for v in vel)
        avg = total_d / len(vel)
        reliability = round(100 * total_d / total_c) if total_c else 0
        for i, (label, value) in enumerate([("Average velocity", f"{avg:.1f}"),
                                            ("Commitment reliability", f"{reliability}%")]):
            y = 1.6 + i * 2.5
            add_rect(slide, 9.5, y, 3.2, 2.2, "gray", name="Card")
            add_rect(slide, 9.5, y, 3.2, 0.07, "blue", name="CardAccent")
            add_text(slide, 9.5, y + 0.45, 3.2, 0.9, value, size=40, bold=True, color="blue",
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="KPI")
            add_text(slide, 9.5, y + 1.5, 3.2, 0.4, label, size=14, align=PP_ALIGN.CENTER,
                     name="KPILabel")

    def next_week_slides(self) -> None:
        rows = [([str(n), i["title"], i["owner"]], None)
                for n, i in enumerate(self.data["next_week"], start=1)]
        self.table_slides("Next Week Plan", "next", ["#", "Planned work", "Owner"],
                          [0.8, 8.9, 2.4], rows, "Plan to be confirmed at sprint planning.")

    def closing_slide(self) -> None:
        d = self.data
        slide = self.hero()
        add_picture(slide, theme.ASSETS_DIR / "acke-logo-white.png", 0.8, 0.8, height=0.9,
                    name="ACKE Logo")
        add_text(slide, 0.8, 2.9, 9.5, 1.2, "Thank you", size=54, bold=True, color="white",
                 name="Title")
        add_text(slide, 0.8, 4.3, 9.5, 0.5, "Questions and blocker decisions welcome", size=22,
                 color="light_blue", name="Subtitle")
        if d.get("contact"):
            add_text(slide, 0.8, 5.2, 9.5, 0.4, d["contact"], size=18, color="white",
                     name="Contact")

    def build(self) -> PresentationType:
        self.title_slide()
        self.summary_slide()
        self.goal_slide()
        self.completed_slides()
        self.in_progress_slides()
        self.blocked_slide()
        self.velocity_slide()
        self.next_week_slides()
        self.closing_slide()
        cp = self.prs.core_properties
        cp.title = (f"ACKE.in Weekly Scrum Status - Sprint {self.data['sprint']} "
                    f"{week_label(self.data)}")
        cp.author = "ACKE.in"
        return self.prs


def build(data: dict[str, Any], out: Path) -> Path:
    errors, _ = validate_status(data)
    if errors:
        raise ValueError("invalid status data:\n  " + "\n  ".join(errors))
    missing = theme.missing_assets()
    if missing:
        raise FileNotFoundError(f"missing assets {missing}; run scripts/generate_assets.py")
    out.parent.mkdir(parents=True, exist_ok=True)
    Deck(data).build().save(out)
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    try:
        data = json.loads(args.data.read_text(encoding="utf-8"))
        print(build(data, args.out))
    except (OSError, ValueError, FileNotFoundError) as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
