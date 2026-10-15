#!/usr/bin/env python3
"""Check a generated deck against the ACKE.in brand and layout rules.

Usage: python validate_deck.py deck.pptx
Exit code 1 on errors. Checks: slide count, logo placement, titles, fonts, palette colors,
bullet limits, shapes inside the slide, and estimated text overflow.
"""
from __future__ import annotations

import argparse
import math
import sys
from pathlib import Path

import theme
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

MIN_SLIDES = 9
MAX_BULLETS = 6
MAX_WORDS = 14
CHAR_WIDTH = 0.52  # average Arial glyph width as a fraction of font size
LINE_HEIGHT = 1.2


def _iter_runs(shape):
    frames = []
    if shape.has_text_frame:
        frames.append(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        frames += [c.text_frame for row in shape.table.rows for c in row.cells]
    for tf in frames:
        for p in tf.paragraphs:
            yield from p.runs


def _estimated_height_pt(shape) -> float:
    width_pt = Emu(shape.width).pt
    total = 0.0
    for p in shape.text_frame.paragraphs:
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18.0)
        chars = sum(len(r.text) for r in p.runs)
        lines = max(1, math.ceil(chars * size * CHAR_WIDTH / max(width_pt, 1)))
        spacing = p.space_after.pt if p.space_after else 0
        total += lines * size * LINE_HEIGHT + spacing
    return total


def check(path: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    prs = Presentation(str(path))
    allowed = {theme.PALETTE[k] for k in theme.PALETTE}
    slides = list(prs.slides)
    if len(slides) < MIN_SLIDES:
        errors.append(f"deck has {len(slides)} slides; expected at least {MIN_SLIDES}")

    for n, slide in enumerate(slides, start=1):
        names = {s.name for s in slide.shapes}
        wanted = "ACKE Logo" if n in (1, len(slides)) else "ACKE Mark"
        if wanted not in names:
            errors.append(f"slide {n}: missing '{wanted}' picture")
        title = next((s for s in slide.shapes if s.name == "Title"), None)
        if title is None or not title.text_frame.text.strip():
            errors.append(f"slide {n}: missing or empty Title")

        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or shape.name != "Background":
                right, bottom = shape.left + shape.width, shape.top + shape.height
                if shape.left < 0 or shape.top < 0 or right > prs.slide_width \
                        or bottom > prs.slide_height:
                    errors.append(f"slide {n}: '{shape.name}' extends beyond the slide")
            for run in _iter_runs(shape):
                if run.font.name != theme.FONT:
                    errors.append(f"slide {n}: '{shape.name}' run not set to {theme.FONT}")
                if run.font.color and run.font.color.type is not None \
                        and str(run.font.color.rgb) not in allowed:
                    errors.append(
                        f"slide {n}: '{shape.name}' uses off-palette {run.font.color.rgb}")
                if "TBD" in run.text:
                    warnings.append(f"slide {n}: placeholder 'TBD' in '{shape.name}'")
            if shape.name == "Bullets":
                paras = shape.text_frame.paragraphs
                if len(paras) > MAX_BULLETS:
                    errors.append(f"slide {n}: {len(paras)} bullets exceeds {MAX_BULLETS}")
                for p in paras:
                    if len("".join(r.text for r in p.runs).split()) > MAX_WORDS:
                        warnings.append(f"slide {n}: bullet longer than {MAX_WORDS} words")
            if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and shape.text_frame.text.strip():
                need, have = _estimated_height_pt(shape), Emu(shape.height).pt
                if need > have * 1.05:
                    errors.append(
                        f"slide {n}: '{shape.name}' text may overflow "
                        f"(needs ~{need:.0f}pt, box {have:.0f}pt)")
    return errors, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("deck", type=Path)
    args = parser.parse_args()
    if not args.deck.exists():
        print(f"ERROR {args.deck} not found", file=sys.stderr)
        return 1
    errors, warnings = check(args.deck)
    for w in warnings:
        print(f"WARN  {w}")
    for e in errors:
        print(f"ERROR {e}")
    print(f"{len(errors)} error(s), {len(warnings)} warning(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
